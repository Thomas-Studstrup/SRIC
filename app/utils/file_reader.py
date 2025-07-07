import os
import json
import docx
import PyPDF2
import openpyxl
import email
from email import policy
import pytesseract
from pdf2image import convert_from_path

def extract_text_from_file(filepath):
    ext = filepath.lower()
    if ext.endswith(".pdf"):
        return extract_text_from_pdf(filepath)
    elif ext.endswith(".docx"):
        return extract_text_from_docx(filepath)
    elif ext.endswith(".xlsx"):
        return extract_text_from_excel(filepath)
    elif ext.endswith(".eml"):
        return extract_text_from_eml(filepath)
    elif ext.endswith(".json"):
        with open(filepath, "r", encoding="utf-8") as f:
            return json.dumps(json.load(f))
    return ""

def extract_text_from_pdf(path):
    try:
        text = ""
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                content = page.extract_text()
                if content:
                    text += content
        if text.strip():
            return text
        else:
            images = convert_from_path(path)
            return " ".join(pytesseract.image_to_string(img) for img in images)
    except Exception as e:
        return f"PDF error: {e}"

def extract_text_from_docx(path):
    try:
        doc = docx.Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        return f"DOCX error: {e}"

def extract_text_from_excel(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.extend([str(cell) for cell in row if cell is not None])
        return " ".join(text)
    except Exception as e:
        return f"Excel error: {e}"

def extract_text_from_eml(path):
    try:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
            body = msg.get_body(preferencelist=('plain',))
            return body.get_content() if body else ""
    except Exception as e:
        return f"EML error: {e}"
